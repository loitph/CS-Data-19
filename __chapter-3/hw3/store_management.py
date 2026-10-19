from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
import os
from copy import copy
from decimal import Decimal

excel_file = "data_nhap_xuat_kho.xlsx"

if os.path.exists(excel_file):
  workbook = load_workbook(excel_file)
  worksheet = workbook.active
else:
  workbook = Workbook()
  worksheet = workbook.active

for row_index, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
  if row_index == 1:
    continue
  quantity, unit_price, import_price = row[3], row[4], row[5]
  if quantity is None or unit_price is None or import_price is None:
    continue
  profit = (Decimal(unit_price) - Decimal(import_price)) * Decimal(quantity)
  type_of_profit = ""
  if profit > 200_000:
    type_of_profit = "Lợi nhuận cao"
  elif 100_000 <= profit <= 200_000:
    type_of_profit = "Lợi nhuận trung bình"
  else:
    type_of_profit = "Lợi nhuận thấp"
  
  cell_profit = worksheet.cell(row=row_index, column=8, value=(
    profit
  ))
  cell_profit._style = copy(worksheet.cell(row=row_index, column=1)._style)
  # format cell #,##0
  cell_profit.number_format = '#,##0'

  cell_type_of_profit = worksheet.cell(row=row_index, column=7, value=(
    type_of_profit
  ))
  cell_type_of_profit._style = copy(worksheet.cell(row=row_index, column=1)._style)

# widen the profit type / profit columns
worksheet.column_dimensions[get_column_letter(7)].width = 22
worksheet.column_dimensions[get_column_letter(8)].width = 15


workbook.save(excel_file)
print(f'Excel file "{excel_file}" updated successfully.')