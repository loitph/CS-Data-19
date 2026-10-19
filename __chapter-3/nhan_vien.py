import os
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from copy import copy
from openpyxl.utils import get_column_letter

excel_file = "nhan_vien.xlsx"

if os.path.exists(excel_file):
  workbook = load_workbook(excel_file)
  worksheet = workbook.active
else:
  workbook = Workbook()
  worksheet = workbook.active
  worksheet.title = 'Nhan_vien'

for row_index, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
  if row_index == 1:
    cell = worksheet.cell(row=row_index, column=9, value="Tổng lương")
    cell._style = copy(worksheet.cell(row=row_index, column=8)._style)
    continue
  cell = worksheet.cell(row=row_index, column=9, value=(int(row[5]) + int(row[6])))
  cell._style = copy(worksheet.cell(row=row_index, column=8)._style)
  # format cell #,##
  cell.number_format = '#,##'

# width 15
worksheet.column_dimensions[get_column_letter(9)].width = 15


workbook.save(excel_file)
print(f'Excel file "{excel_file}" updated successfully.')