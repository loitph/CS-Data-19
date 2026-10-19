import os
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from copy import copy


excel_file = "ban_hang.xlsx"

HEADERS = [
  "Ngày bán",
  "Mã đơn hàng",
  "Sản phẩm",
  "Số lượng",
  "Đơn giá",
  "Khu vực",
  "Nhân viên",
  "Tổng tiền"
]

if os.path.exists(excel_file):
  workbook = load_workbook(excel_file)
  worksheet = workbook.active
else:
  workbook = Workbook()
  worksheet = workbook.active
  worksheet.title = 'Ban_hang'
  worksheet.append(HEADERS)

for row_index, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
  if row_index == 1:
    cell = worksheet.cell(row=row_index, column=8, value="Tổng tiền")
    cell._style = copy(worksheet.cell(row=row_index, column=7)._style)
    continue
  cell = worksheet.cell(row=row_index, column=8, value=row[3] * row[4])
  cell._style = copy(worksheet.cell(row=row_index, column=7)._style)
  # cell.font = copy(worksheet.cell(row=row_index, column=7).font)
  # cell.font = Font(bold=True)

workbook.save(excel_file)
print(f'Excel file "{excel_file}" updated successfully.')
