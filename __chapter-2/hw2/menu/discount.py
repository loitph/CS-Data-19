import os
import re
from decimal import Decimal, InvalidOperation
from openpyxl import Workbook, load_workbook

class Order:
  def __init__(self, name: str, price: Decimal, quantity: int, rate: Decimal):
    self.name = name
    self.price = price
    self.quantity = quantity
    self.subtotal = price * quantity
    self.rate = rate
    self.discount = self.subtotal * rate / 100
    self.total = self.subtotal - self.discount
    self.extra_discount = calc_extra_discount(self.total)
    self.total_pay = self.total - self.extra_discount

  def __str__(self) -> str:
    return f"{self.name}, {self.price}, {self.quantity}, {self.subtotal}, {self.rate}, {self.discount}, {self.total}"

  def __to_row__(self) -> list[str | int | Decimal]:
    return [self.name, self.price, self.quantity, self.subtotal, self.rate, self.discount, self.total]

  def get_row(self) -> list[str | int | Decimal]:
    return self.__to_row__()


menu_title = "\n--- Menu ---"
opt_1_info = "1. Enter the product's name, price per unit, quantity"
opt_2_print = "2. Print the bill"
opt_3_save = "3. Save the bill to Excel"
opt_4_exit = "4. Exit"
menu_end = "--- #### ---\n"

msg_choice = "Enter your choice: "
msg_err_choice = "Invalid choice. Please enter a number between 1 and 4.\n"

msg_product_name = "Enter the product's name: "
msg_err_name = "Invalid name. Letters, digits and spaces only.\n"
msg_price = "Enter the price per unit (VND): "
msg_err_price = "Invalid input. Please enter a valid price greater than 0.\n"
msg_quantity = "Enter the quantity: "
msg_err_quantity = "Invalid input. Please enter a whole number greater than 0.\n"

goodbye = "Goodbye!"

PRODUCT_PATTERN = re.compile(r"[^\W_]+(?: [^\W_]+)*")

EXTRA_DISCOUNT_THRESHOLD = Decimal(10_000_000)
EXTRA_DISCOUNT = Decimal(200_000)

lst_orders: list[Order] = []
saved_count = 0

def safe_input_name(s: str) -> str | None:
  s = " ".join(s.split())
  if PRODUCT_PATTERN.fullmatch(s):
    return s
  return None

def safe_parse_int(s: str) -> int | None:
  try:
    return int(s)
  except ValueError:
    return None

def safe_parse_decimal(s: str) -> Decimal | None:
  try:
    d = Decimal(s.strip().replace("_", "").replace(",", ""))
  except InvalidOperation:
    return None
  if not d.is_finite():
    return None
  return d

def choice_input(prompt: str) -> int:
  while True:
    choice = safe_parse_int(input(prompt))
    if choice is not None and 1 <= choice <= 4:
      return choice
    print(msg_err_choice)

def name_input(prompt: str) -> str:
  while True:
    name = safe_input_name(input(prompt))
    if name is not None:
      return name
    print(msg_err_name)

def price_input(prompt: str) -> Decimal:
  while True:
    price = safe_parse_decimal(input(prompt))
    if price is not None and price > 0:
      return price
    print(msg_err_price)

def quantity_input(prompt: str) -> int:
  while True:
    quantity = safe_parse_int(input(prompt))
    if quantity is not None and quantity > 0:
      return quantity
    print(msg_err_quantity)

def calc_rate(quantity: int) -> Decimal:
  # quantity < 10 -> 0%
  # 10 <= quantity < 20 -> 5%
  # 20 <= quantity < 50 -> 10%
  # quantity >= 50 -> 15%

  TIERS = [
    (Decimal(10), Decimal(0)),
    (Decimal(20), Decimal(5)),
    (Decimal(50), Decimal(10)),
    (Decimal('Infinity'), Decimal(15)),
  ]

  for limit, rate in TIERS:
    if quantity < limit:
      return rate
  return TIERS[-1][1]

def calc_extra_discount(total: Decimal) -> Decimal:
  return EXTRA_DISCOUNT if total > EXTRA_DISCOUNT_THRESHOLD else Decimal(0)

def build_order(name: str, price: Decimal, quantity: int) -> Order:
  return Order(name, price, quantity, calc_rate(quantity))

def fmt_vnd(amount: Decimal) -> str:
  return f"{amount:,.0f} VND"

def print_bill() -> None:
  if not lst_orders:
    print('No data to print.')
    return

  for order in lst_orders:
    print("\n========= Bill =========")
    print(f"Product's name: {order.name}")
    print(f"Price per unit: {fmt_vnd(order.price)}")
    print(f"Quantity: {order.quantity}")
    print(f"Total before discount: {fmt_vnd(order.subtotal)}")
    print(f"Discount (%): {order.rate}")
    print(f"Discount amount: {fmt_vnd(order.discount)}")
    print(f"Total: {fmt_vnd(order.total)}")
    if order.extra_discount:
      print(f"Extra discount (total > {fmt_vnd(EXTRA_DISCOUNT_THRESHOLD)}): -{fmt_vnd(order.extra_discount)}")
    print(f"=> Total should pay: {fmt_vnd(order.total_pay)}")
    print("======================================")

def save_to_excel() -> None:
  global saved_count
  excel_file = 'discount.xlsx'

  new_orders = lst_orders[saved_count:]
  if not new_orders:
    print('No new data to save.')
    return

  if os.path.exists(excel_file):
    workbook = load_workbook(excel_file)
    worksheet = workbook.active
  else:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = 'Discount Report'
    worksheet.append([
      'No.',
      "Product's name",
      'Price per unit',
      'Quantity',
      'Total before discount',
      'Discount (%)',
      'Discount amount',
      'Total'
    ])

  start_no = worksheet.max_row
  for i, order in enumerate(new_orders, start=start_no):
    worksheet.append([i, *order.get_row()])

  workbook.save(excel_file)
  print(f'Saved {len(new_orders)} new row(s) to "{excel_file}".')

  saved_count = len(lst_orders)


def main() -> None:
  os.system('cls' if os.name == 'nt' else 'clear')
  while True:
    print(menu_title)
    print(opt_1_info)
    print(opt_2_print)
    print(opt_3_save)
    print(opt_4_exit)
    print(menu_end)

    choice = choice_input(msg_choice)

    if choice == 4:
      print(goodbye)
      break
    elif choice == 1:
      name = name_input(msg_product_name)
      price = price_input(msg_price)
      quantity = quantity_input(msg_quantity)
      lst_orders.append(build_order(name, price, quantity))
    elif choice == 2:
      print_bill()
    elif choice == 3:
      save_to_excel()
    else:
      print(msg_err_choice)

if __name__ == "__main__":
  main()
