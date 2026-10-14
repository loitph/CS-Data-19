import os
import re
from decimal import Decimal, InvalidOperation
from openpyxl import Workbook, load_workbook

class Clerk:
  def __init__(self, name: str, sales: Decimal, rate: Decimal, commission: Decimal, bonus: Decimal, note: str):
    self.name = name
    self.sales = sales
    self.rate = rate
    self.commission = commission
    self.bonus = bonus
    self.total = commission + bonus
    self.note = note

  def __str__(self) -> str:
    return f"{self.name}, {self.sales}, {self.rate}, {self.commission}, {self.bonus}, {self.total}, {self.note}"

  def __to_row__(self) -> list[str | Decimal]:
    return [self.name, self.sales, self.rate, self.commission, self.bonus, self.total, self.note]

  def get_row(self) -> list[str | Decimal]:
    return self.__to_row__()


menu_title = "\n--- Menu ---"
opt_1_info = "1. Enter the clerk's name, sales volume"
opt_2_print = "2. Print the report"
opt_3_save = "3. Save the report to Excel"
opt_4_exit = "4. Exit"
menu_end = "--- #### ---\n"

msg_choice = "Enter your choice: "
msg_err_choice = "Invalid choice. Please enter a number between 1 and 4.\n"

msg_clerk_name = "Enter the clerk's name: "
msg_err_name = "Invalid name. Letters and spaces only.\n"
msg_sales = "Enter the sales volume (VND): "
msg_err_sales = "Invalid input. Please enter a valid sales volume.\n"

msg_low_performance = "Low performance, please enhance"

goodbye = "Goodbye!"

WORDS_PATTERN = re.compile(r"[^\W\d_]+(?: [^\W\d_]+)*")

BONUS_THRESHOLD = Decimal(70_000_000)
BONUS_ALLOWANCE = Decimal(2_000_000)
LOW_PERFORMANCE_THRESHOLD = Decimal(5_000_000)

lst_clerks: list[Clerk] = []
saved_count = 0

def safe_input_name(s: str) -> str | None:
  s = " ".join(s.split())
  if WORDS_PATTERN.fullmatch(s):
    return s
  return None

def safe_parse_int(s: str) -> int | None:
  try:
    return int(s)
  except ValueError:
    return None

def safe_parse_decimal(s: str) -> Decimal | None:
  try:
    d = Decimal(s.strip().replace("_", "").replace(",", ""))
  except InvalidOperation:
    return None
  if not d.is_finite():
    return None
  return d

def choice_input(prompt: str) -> int:
  while True:
    choice = safe_parse_int(input(prompt))
    if choice is not None and 1 <= choice <= 4:
      return choice
    print(msg_err_choice)

def name_input(prompt: str) -> str:
  while True:
    name = safe_input_name(input(prompt))
    if name is not None:
      return name
    print(msg_err_name)

def sales_input(prompt: str) -> Decimal:
  while True:
    sales = safe_parse_decimal(input(prompt))
    if sales is not None and sales >= 0:
      return sales
    print(msg_err_sales)

def calc_rate(sales: Decimal) -> Decimal:
  # sales < 10_000_000 -> 0%
  # 10_000_000 <= sales < 30_000_000 -> 5%
  # 30_000_000 <= sales < 50_000_000 -> 10%
  # sales >= 50_000_000 -> 15%

  TIERS = [
    (Decimal(10_000_000), Decimal(0)),
    (Decimal(30_000_000), Decimal(5)),
    (Decimal(50_000_000), Decimal(10)),
    (Decimal('Infinity'), Decimal(15)),
  ]

  for limit, rate in TIERS:
    if sales < limit:
      return rate
  return TIERS[-1][1]

def calc_bonus(sales: Decimal) -> Decimal:
  return BONUS_ALLOWANCE if sales >= BONUS_THRESHOLD else Decimal(0)

def calc_note(sales: Decimal) -> str:
  return msg_low_performance if sales < LOW_PERFORMANCE_THRESHOLD else ""

def build_clerk(name: str, sales: Decimal) -> Clerk:
  rate = calc_rate(sales)
  commission = sales * rate / 100
  return Clerk(name, sales, rate, commission, calc_bonus(sales), calc_note(sales))

def fmt_vnd(amount: Decimal) -> str:
  return f"{amount:,.0f} VND"

def print_report() -> None:
  if not lst_clerks:
    print('No data to print.')
    return

  for clerk in lst_clerks:
    print("\n===== BÁO CÁO HOA HỒNG NHÂN VIÊN =====")
    print(f"Name: {clerk.name}")
    print(f"Sales volume: {fmt_vnd(clerk.sales)}")
    print(f"Commission rate (%): {clerk.rate}")
    print(f"Commission: {fmt_vnd(clerk.commission)}")
    print(f"Bonus allowance: {fmt_vnd(clerk.bonus)}")
    if clerk.note:
      print(f"Warning: {clerk.note}")
    print(f"=> Total Income: {fmt_vnd(clerk.total)}")
    print("======================================")

def save_to_excel() -> None:
  global saved_count
  excel_file = 'sales-commission.xlsx'

  new_clerks = lst_clerks[saved_count:]
  if not new_clerks:
    print('No new data to save.')
    return

  if os.path.exists(excel_file):
    workbook = load_workbook(excel_file)
    worksheet = workbook.active
  else:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = 'Sales Commission Report'
    worksheet.append([
      'No.',
      'Name',
      'Sales Volume',
      'Commission rate (%)',
      'Commission (VND)',
      'Bonus Allowance',
      'Total Income',
      'Note'
    ])

  start_no = worksheet.max_row
  for i, clerk in enumerate(new_clerks, start=start_no):
    worksheet.append([i, *clerk.get_row()])

  workbook.save(excel_file)
  print(f'Saved {len(new_clerks)} new row(s) to "{excel_file}".')

  saved_count = len(lst_clerks)


def main() -> None:
  os.system('cls' if os.name == 'nt' else 'clear')
  while True:
    print(menu_title)
    print(opt_1_info)
    print(opt_2_print)
    print(opt_3_save)
    print(opt_4_exit)
    print(menu_end)

    choice = choice_input(msg_choice)

    if choice == 4:
      print(goodbye)
      break
    elif choice == 1:
      name = name_input(msg_clerk_name)
      sales = sales_input(msg_sales)
      lst_clerks.append(build_clerk(name, sales))
    elif choice == 2:
      print_report()
    elif choice == 3:
      save_to_excel()
    else:
      print(msg_err_choice)

if __name__ == "__main__":
  main()
